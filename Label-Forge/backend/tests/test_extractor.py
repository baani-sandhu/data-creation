import pytest

from app.services.extractor import extract


def test_pdf_extraction_returns_text(sample_pdf_bytes):
    text = extract(sample_pdf_bytes, "pdf")
    assert isinstance(text, str)
    assert len(text) > 100
    assert text.strip() != ""


def test_txt_extraction(sample_txt_bytes):
    text = extract(sample_txt_bytes, "txt")
    assert text == sample_txt_bytes.decode("utf-8")


def test_md_extraction():
    content = b"# Header\n\nSome paragraph text here."
    text = extract(content, "md")
    assert "Header" in text
    assert "paragraph" in text


def test_csv_extraction():
    content = b"name,age\nAlice,30\nBob,25"
    text = extract(content, "csv")
    assert "Alice" in text
    assert "Bob" in text


def test_pdf_extraction_returns_nonempty_for_real_pdf(sample_pdf_bytes):
    text = extract(sample_pdf_bytes, "pdf")
    word_count = len(text.split())
    assert word_count > 50, f"Expected >50 words, got {word_count}"


def test_docx_extraction():
    from docx import Document as DocxDocument
    import io
    doc = DocxDocument()
    doc.add_heading("Test Heading", level=1)
    doc.add_paragraph("This is a paragraph.")
    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    text = extract(file_bytes, "docx")
    assert "Test Heading" in text
    assert "paragraph" in text


def test_xlsx_extraction():
    import openpyxl
    import io
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Age", "City"])
    ws.append(["Alice", 30, "Delhi"])
    ws.append(["Bob", 25, "Mumbai"])
    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    text = extract(file_bytes, "xlsx")
    assert "Alice" in text
    assert "Delhi" in text
    assert "Bob" in text


def test_pptx_extraction():
    from pptx import Presentation
    from pptx.util import Inches
    import io
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Test Slide"
    slide.placeholders[1].text = "Bullet point content"
    buf = io.BytesIO()
    prs.save(buf)
    file_bytes = buf.getvalue()

    text = extract(file_bytes, "pptx")
    assert "Test Slide" in text
    assert "Bullet point content" in text


def test_unsupported_extension_raises():
    with pytest.raises(ValueError, match="Unsupported"):
        extract(b"some bytes", "xyz")


def test_mixed_format_extraction():
    formats_and_content = [
        (b"Hello world this is plain text", "txt"),
        (b"name,age\nAlice,30", "csv"),
    ]
    for content, fmt in formats_and_content:
        result = extract(content, fmt)
        assert len(result.strip()) > 0
