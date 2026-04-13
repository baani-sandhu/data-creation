import pdfplumber
import fitz
from pptx import Presentation
from docx import Document as DocxDocument
import openpyxl
import xlrd
import io

def _extract_pdf(file_bytes: bytes) -> str:
    text = _try_pymupdf(file_bytes)
    if not text.strip():
        text = _try_pdfplumber(file_bytes)
    return text

def _try_pdfplumber(file_bytes: bytes) -> str:
    pages = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                pages.append(t.strip())
    return "\n\n".join(pages)

def _try_pymupdf(file_bytes: bytes) -> str:
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    try:
        pages = []
        for page_index in range(doc.page_count):
            t = doc.load_page(page_index).get_text()
            if t.strip():
                pages.append(t.strip())

        return "\n\n".join(pages)
    finally:
        doc.close()

def _extract_text(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore")


def _extract_pptx(file_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(file_bytes))
    slides_text = []
    for slide in prs.slides:
        slide_lines = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                line = para.text.strip()
                if line:
                    slide_lines.append(line)
        if slide_lines:
            slides_text.append("\n".join(slide_lines))
    return "\n\n".join(slides_text)


def _extract_docx(file_bytes: bytes) -> str:
    doc = DocxDocument(io.BytesIO(file_bytes))
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = ""
        if para.style is not None and para.style.name is not None:
            style_name = para.style.name
        if "Heading" in style_name:
            lines.append(f"\n{text}\n")
        else:
            lines.append(text)
    return "\n".join(lines)


def _extract_xlsx(file_bytes: bytes) -> str:
    wb = openpyxl.load_workbook(
        io.BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )
    all_rows = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                all_rows.append(", ".join(cells))
    return "\n".join(all_rows)


def _extract_xls(file_bytes: bytes) -> str:
    wb = xlrd.open_workbook(file_contents=file_bytes)
    all_rows = []
    for sheet in wb.sheets():
        for row_idx in range(sheet.nrows):
            cells = []
            for col_idx in range(sheet.ncols):
                cell_value = sheet.cell_value(row_idx, col_idx)
                if cell_value != "":
                    cells.append(str(cell_value))
            if cells:
                all_rows.append(", ".join(cells))
    return "\n".join(all_rows)


SUPPORTED_TYPES = {
    "pdf": _extract_pdf,
    "txt": _extract_text,
    "md": _extract_text,
    "csv": _extract_text,
    "pptx": _extract_pptx,
    "docx": _extract_docx,
    "xlsx": _extract_xlsx,
    "xls": _extract_xls,
}


def extract(file_bytes: bytes, file_ext: str) -> str:
    file_ext = file_ext.lower().lstrip(".")
    if file_ext not in SUPPORTED_TYPES:
        raise ValueError(f"Unsupported file type: .{file_ext}")
    return SUPPORTED_TYPES[file_ext](file_bytes)
